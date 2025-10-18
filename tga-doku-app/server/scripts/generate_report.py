#!/usr/bin/env python3
"""
Generate Word document report from project data
"""
import sys
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from datetime import datetime

def create_report(project_data):
    """Create a Word document report"""
    doc = Document()
    
    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)
    
    project = project_data['project']
    findings = project_data['findings']
    floor_plans = project_data['floorPlans']
    photos = project_data['photos']
    
    # Title Page
    title = doc.add_heading(project['name'], 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    subtitle = doc.add_paragraph('Begehungs- und Übergabedokumentation')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.size = Pt(14)
    subtitle.runs[0].font.color.rgb = RGBColor(100, 100, 100)
    
    doc.add_paragraph()
    
    # Project Info Table
    doc.add_heading('Grunddaten des Projekts', 1)
    table = doc.add_table(rows=4, cols=2)
    table.style = 'Light Grid Accent 1'
    
    table.rows[0].cells[0].text = 'Projekt:'
    table.rows[0].cells[1].text = project['name']
    table.rows[1].cells[0].text = 'Standort:'
    table.rows[1].cells[1].text = project.get('location', 'N/A')
    table.rows[2].cells[0].text = 'Datum:'
    table.rows[2].cells[1].text = datetime.now().strftime('%d.%m.%Y')
    table.rows[3].cells[0].text = 'Status:'
    table.rows[3].cells[1].text = project.get('status', 'N/A')
    
    doc.add_page_break()
    
    # Findings Section
    doc.add_heading('Feststellungen', 1)
    
    # Group findings by category
    findings_by_category = {}
    for finding in findings:
        category = finding['category']
        if category not in findings_by_category:
            findings_by_category[category] = []
        findings_by_category[category].append(finding)
    
    # Add findings by category
    for category, category_findings in findings_by_category.items():
        doc.add_heading(category.upper(), 2)
        
        for finding in category_findings:
            doc.add_heading(finding['title'], 3)
            
            if finding.get('description'):
                p = doc.add_paragraph()
                p.add_run('Beschreibung: ').bold = True
                p.add_run(finding['description'])
            
            if finding.get('recommendation'):
                p = doc.add_paragraph()
                p.add_run('Empfehlung: ').bold = True
                p.add_run(finding['recommendation'])
            
            if finding.get('responsibility') or finding.get('deadline'):
                p = doc.add_paragraph()
                if finding.get('responsibility'):
                    p.add_run(f"Verantwortlich: {finding['responsibility']}  ")
                if finding.get('deadline'):
                    p.add_run(f"Frist: {finding['deadline']}")
            
            doc.add_paragraph()  # Spacing
    
    doc.add_page_break()
    
    # Summary Table
    doc.add_heading('Zusammenfassung und Maßnahmenübersicht', 1)
    
    if findings:
        summary_table = doc.add_table(rows=len(findings) + 1, cols=5)
        summary_table.style = 'Light Grid Accent 1'
        
        # Header
        header_cells = summary_table.rows[0].cells
        header_cells[0].text = 'Nr.'
        header_cells[1].text = 'Feststellung'
        header_cells[2].text = 'Empfehlung'
        header_cells[3].text = 'Verantwortlich'
        header_cells[4].text = 'Frist'
        
        # Data
        for idx, finding in enumerate(findings, 1):
            row_cells = summary_table.rows[idx].cells
            row_cells[0].text = str(idx)
            row_cells[1].text = finding['title']
            row_cells[2].text = finding.get('recommendation', '-')[:50]
            row_cells[3].text = finding.get('responsibility', '-')
            row_cells[4].text = finding.get('deadline', '-')
    
    doc.add_page_break()
    
    # Next Steps
    doc.add_heading('Weiteres Vorgehen', 1)
    doc.add_paragraph('Die dokumentierten Feststellungen sind zu prüfen und die empfohlenen Maßnahmen umzusetzen.')
    doc.add_paragraph('Eine Nachbegehung wird empfohlen, um die Umsetzung der Maßnahmen zu verifizieren.')
    
    return doc

def main():
    if len(sys.argv) < 2:
        print("Usage: generate_report.py <project_data_json>")
        sys.exit(1)
    
    project_data = json.loads(sys.argv[1])
    doc = create_report(project_data)
    
    output_path = f"/tmp/report_{project_data['project']['id']}.docx"
    doc.save(output_path)
    print(output_path)

if __name__ == "__main__":
    main()

